# 網頁端對端測試：用 update.mjs 在模擬世界產出的 returns.json 渲染 index.html
# 用法：python3 test/page_test.py <含 data/returns.json 的資料夾> <輸出資料夾> <xlsx.full.min.js 路徑>
import json, sys, threading, http.server, functools, shutil, os
from playwright.sync_api import sync_playwright
import openpyxl

data_dir, out_dir, xlsx_js = sys.argv[1], sys.argv[2], sys.argv[3]
root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
site = os.path.join(out_dir, 'site')
shutil.rmtree(site, ignore_errors=True)
os.makedirs(os.path.join(site, 'data'))
shutil.copy(os.path.join(root, 'index.html'), site)
shutil.copy(os.path.join(data_dir, 'data', 'returns.json'), os.path.join(site, 'data'))

handler = functools.partial(http.server.SimpleHTTPRequestHandler, directory=site)
handler.log_message = lambda *a: None
srv = http.server.ThreadingHTTPServer(('127.0.0.1', 0), handler)
threading.Thread(target=srv.serve_forever, daemon=True).start()
url = f'http://127.0.0.1:{srv.server_port}/index.html'

passed = failed = 0
def ok(c, m):
    global passed, failed
    if c: passed += 1; print('  ✓', m)
    else: failed += 1; print('  ✗', m)

with sync_playwright() as p:
    b = p.chromium.launch()
    ctx = b.new_context(viewport={'width': 1440, 'height': 1100}, accept_downloads=True)
    ctx.route('https://fonts.googleapis.com/**', lambda r: r.abort())
    ctx.route('https://cdnjs.cloudflare.com/**', lambda r: r.fulfill(path=xlsx_js, content_type='application/javascript'))
    pg = ctx.new_page()
    errors = []
    pg.on('pageerror', lambda e: errors.append(str(e)))
    pg.goto(url)
    pg.wait_for_selector('.drow')
    ok(not errors, f'沒有 JS 錯誤 {errors}')
    tabs = pg.locator('.tab').all_inner_texts()
    print('   年度分頁：', tabs)
    ok(len(tabs) == 2 and '今年' in tabs[0], '有今年與 2025 全年兩個分頁')
    heads = pg.locator('.hrow .hcell').all_inner_texts()
    print('   表頭：', [h.replace('\n', ' ') for h in heads])
    pg.screenshot(path=os.path.join(out_dir, 'desktop.png'), full_page=True)

    def row_text(code):
        return pg.locator('.drow', has=pg.locator('.code', has_text=code)).first.inner_text().replace('\n', ' | ')
    for code in ['0050', '00631L', '00415A', '00981A', '00699']:
        if pg.locator('.code', has_text=code).count():
            print(f'   {code}：', row_text(code))

    # 2025 分頁：0050 6 月分割後不能出現 −75%
    pg.click('.tab >> nth=1')
    r0050 = row_text('0050')
    print('   2025 0050：', r0050)
    ok('待確認' not in r0050 and '−7' not in r0050, '0050 2025/6 分割已換算，沒有 −75% 也沒有待確認')
    r981 = row_text('00981A')
    ok('掛牌月' in r981 and '掛牌以來' in r981, '00981A 2025 標示掛牌月與掛牌以來')
    r699 = row_text('00699')
    ok('已下市' in r699, '00699 標已下市')
    r777 = row_text('00777')
    ok('缺資料' not in r777, '00777 月底停牌的月份有數字')
    pg.screenshot(path=os.path.join(out_dir, 'desktop-2025.png'), full_page=True)

    # 篩選：大型權值＋股票型＋國內
    pg.click('.tab >> nth=0')
    pg.click('.chip[data-k="theme"][data-v="大型權值"]')
    pg.click('.chip[data-k="type"][data-v="股票型"]')
    pg.click('.chip[data-k="region"][data-v="國內"]')
    codes = pg.locator('.drow .code').all_inner_texts()
    print('   大型權值＋股票型＋國內：', codes)
    ok(set(codes) == {'0050', '006208'}, '只剩 0050、006208（正2、反1、S&P500、主動式被濾掉）')
    pg.click('.chip[data-k="type"][data-v="多資產"]')
    ok(pg.locator('.empty').count() == 1, '沒有結果時顯示提示')
    for k in ['theme', 'type', 'region']:
        pg.click(f'.chip[data-k="{k}"][data-v="全部"]')
    pg.fill('#q', '00631')
    ok(pg.locator('.drow').count() == 1, '搜尋代號')
    pg.fill('#q', '')

    # 排序：點今年以來兩次 → 由小到大
    vals_desc = pg.locator('.ytdcell .v').all_inner_texts()
    pg.click('.hcell.ytd')
    vals_asc = pg.locator('.ytdcell .v').all_inner_texts()
    def num(s):
        s = s.replace('−', '-').replace('+', '').replace('%', '')
        try: return float(s)
        except: return None
    d = [num(x) for x in vals_desc if num(x) is not None]
    a = [num(x) for x in vals_asc if num(x) is not None]
    ok(d == sorted(d, reverse=True), '預設依今年以來由高到低')
    ok(a == sorted(a), '再點一次由低到高')

    # 篩選記住
    pg.click('.chip[data-k="market"][data-v="上櫃"]')
    pg.reload(); pg.wait_for_selector('.drow')
    ok(pg.locator('.chip[data-k="market"][data-v="上櫃"]').get_attribute('aria-pressed') == 'true', '重新整理後篩選仍保留')
    tp = pg.locator('.drow .code').all_inner_texts()
    print('   上櫃：', tp)
    ok('00888' in tp and '依名稱' in row_text('00888'), '上櫃依名稱判斷的高股息有註明')

    # Excel
    pg.click('.chip[data-k="market"][data-v="全部"]')
    with pg.expect_download() as dl:
        pg.click('#xlsx')
    f = os.path.join(out_dir, dl.value.suggested_filename)
    dl.value.save_as(f)
    print('   下載檔名：', dl.value.suggested_filename)
    wb = openpyxl.load_workbook(f)
    ok(wb.sheetnames == ['2026', '2025'], f'工作表 {wb.sheetnames}')
    ws = wb['2026']
    hdr = [c.value for c in ws[1]]
    print('   2026 表頭：', hdr)
    first = [c.value for c in ws[2]]
    print('   2026 第一列：', first)
    num_cell = ws.cell(row=2, column=7)
    ok(isinstance(num_cell.value, (int, float)) or num_cell.value is None, '報酬率是數字')
    ok(any(ws.cell(row=r, column=7).number_format == '0.00%' for r in range(2, ws.max_row + 1)), '百分比格式 0.00%')
    ok(hdr[-2] == '今年以來' and hdr[-1] == '備註', '最後兩欄是今年以來與備註')
    ws25 = wb['2025']
    ok(ws25.cell(row=1, column=7 + 11).value == '12月' and ws25.cell(row=1, column=19).value == '2025 全年', '2025 表 12 個月＋全年')

    # 手機版
    m = b.new_context(viewport={'width': 390, 'height': 844}, device_scale_factor=2)
    m.route('https://fonts.googleapis.com/**', lambda r: r.abort())
    mp = m.new_page()
    mp.goto(url); mp.wait_for_selector('.drow')
    sw = mp.evaluate('document.documentElement.scrollWidth')
    ok(sw <= 390, f'手機版頁面沒有橫向溢出（scrollWidth={sw}）')
    mp.screenshot(path=os.path.join(out_dir, 'mobile.png'), full_page=False)
    mp.evaluate("document.querySelector('.tablebox').scrollLeft = 400")
    mp.locator('.tablebox').screenshot(path=os.path.join(out_dir, 'mobile-scrolled.png'))
    b.close()
srv.shutdown()
print(f'\n結果：{passed} 通過、{failed} 失敗')
sys.exit(1 if failed else 0)
